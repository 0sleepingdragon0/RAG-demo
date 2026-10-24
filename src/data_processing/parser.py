"""
文档解析器：支持 PDF / Word / Markdown，调用 MinerU 进行地质文档解析。
"""
import json
import logging
from pathlib import Path
from dataclasses import dataclass, field
from typing import Optional

from config.settings import RAW_DIR, PARSED_DIR

logger = logging.getLogger(__name__)


@dataclass
class ParsedDocument:
    """解析后的文档结构"""
    source_file: str
    doc_type: str  # "勘查报告" / "论文" / "地质图" / "结构化数据"
    title: str = ""
    text: str = ""
    tables: list = field(default_factory=list)  # [{"caption":..., "content":...}]
    figures: list = field(default_factory=list)  # [{"caption":..., "page":...}]
    page_count: int = 0
    metadata: dict = field(default_factory=dict)


class DocumentParser:
    """
    文档解析器。

    支持的策略：
    - PDF（含扫描图）：优先调用 MinerU 解析为 Markdown，回退 pymupdf 纯文本提取。
    - Word (.docx)：用 python-docx 提取段落 + 表格。
    - Markdown / TXT：直接读取。
    - Excel / CSV：用 pandas 转为结构化文本块。

    输出：将每个文档解析结果保存为 PARSED_DIR 下的 JSON 文件。
    """

    def __init__(self, raw_dir: Optional[Path] = None, parsed_dir: Optional[Path] = None):
        self.raw_dir = raw_dir or RAW_DIR
        self.parsed_dir = parsed_dir or PARSED_DIR
        self.parsed_dir.mkdir(parents=True, exist_ok=True)

    def parse_all(self) -> list[ParsedDocument]:
        """批量解析 raw_dir 下所有支持的文件"""
        docs = []
        for file_path in sorted(self.raw_dir.iterdir()):
            if file_path.suffix.lower() in (".pdf", ".docx", ".doc", ".md", ".txt", ".xlsx", ".csv"):
                try:
                    doc = self.parse_file(file_path)
                    docs.append(doc)
                    logger.info(f"解析完成: {file_path.name} ({len(doc.text)} 字符)")
                except Exception as e:
                    logger.error(f"解析失败: {file_path.name} — {e}")
        logger.info(f"共解析 {len(docs)} 个文档")
        return docs

    def parse_file(self, file_path: Path) -> ParsedDocument:
        """根据文件类型选择解析策略"""
        suffix = file_path.suffix.lower()
        if suffix == ".pdf":
            return self._parse_pdf(file_path)
        elif suffix in (".docx", ".doc"):
            return self._parse_docx(file_path)
        elif suffix in (".md", ".txt"):
            return self._parse_text(file_path)
        elif suffix in (".xlsx", ".csv"):
            return self._parse_tabular(file_path)
        else:
            raise ValueError(f"不支持的文件类型: {suffix}")

    def _parse_pdf(self, file_path: Path) -> ParsedDocument:
        """
        PDF 解析：优先 MinerU，回退 pymupdf。
        MinerU 能输出结构化 Markdown + 表格，适合地质报告。
        """
        # 尝试 MinerU
        mineru_output = self._try_mineru(file_path)
        if mineru_output:
            return mineru_output

        # 回退 pymupdf
        return self._parse_pdf_pymupdf(file_path)

    def _try_mineru(self, file_path: Path) -> Optional[ParsedDocument]:
        """
        尝试调用 MinerU 解析。

        MinerU 2.x/4.x 的命令行是 `mineru`；早期版本(1.x)叫 `magic-pdf`。
        两者都试一次，任一成功即用。
        """
        for cli in ("mineru", "magic-pdf"):
            doc = self._run_mineru_cli(cli, file_path)
            if doc is not None:
                return doc
        return None

    def _run_mineru_cli(self, cli: str, file_path: Path) -> Optional[ParsedDocument]:
        """执行单个 MinerU 命令并收集输出"""
        import subprocess

        output_dir = self.parsed_dir / file_path.stem / cli
        output_dir.mkdir(parents=True, exist_ok=True)

        try:
            result = subprocess.run(
                [cli, "-p", str(file_path), "-o", str(output_dir)],
                capture_output=True, text=True, timeout=1800,
            )
            if result.returncode != 0:
                logger.debug(f"{cli} 退出码 {result.returncode}: {result.stderr[:200]}")
                return None
        except FileNotFoundError:
            logger.debug(f"未安装 {cli}，跳过")
            return None
        except subprocess.TimeoutExpired:
            logger.warning(f"{cli} 解析超时: {file_path.name}")
            return None
        except Exception as e:
            logger.warning(f"{cli} 调用异常: {e}")
            return None

        # 输出目录结构随版本变化，递归找主 Markdown 比较稳妥
        md_candidates = sorted(
            (p for p in output_dir.rglob("*.md") if "images" not in p.parts),
            key=lambda p: p.stat().st_size, reverse=True,
        )
        if not md_candidates:
            logger.warning(f"{cli} 运行成功但未找到输出 Markdown")
            return None

        md_file = md_candidates[0]
        text = md_file.read_text(encoding="utf-8", errors="replace")

        # 表格文件（若有）
        tables: list = []
        for tf in output_dir.rglob("tables.json"):
            try:
                tables = json.loads(tf.read_text(encoding="utf-8"))
                break
            except Exception:
                continue

        logger.info(f"MinerU({cli}) 解析成功: {file_path.name} → {md_file.name}")
        return ParsedDocument(
            source_file=file_path.name,
            doc_type=self._infer_doc_type(file_path.name),
            title=file_path.stem,
            text=text,
            tables=tables,
            metadata={"parser": f"MinerU/{cli}", "md_path": str(md_file)},
        )

    def _parse_pdf_pymupdf(self, file_path: Path) -> ParsedDocument:
        """pymupdf 纯文本提取（回退方案）"""
        import fitz  # pymupdf

        doc = fitz.open(str(file_path))
        pages_text = []
        for page in doc:
            pages_text.append(page.get_text())
        full_text = "\n\n".join(pages_text)

        return ParsedDocument(
            source_file=file_path.name,
            doc_type=self._infer_doc_type(file_path.name),
            title=file_path.stem,
            text=full_text,
            page_count=len(doc),
            metadata={"parser": "pymupdf"},
        )

    def _parse_docx(self, file_path: Path) -> ParsedDocument:
        """Word 文档解析：提取段落和表格"""
        from docx import Document

        doc = Document(str(file_path))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        text = "\n".join(paragraphs)

        # 提取表格
        tables = []
        for i, table in enumerate(doc.tables):
            rows = []
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                rows.append(cells)
            tables.append({"table_index": i, "rows": rows})

        return ParsedDocument(
            source_file=file_path.name,
            doc_type=self._infer_doc_type(file_path.name),
            title=file_path.stem,
            text=text,
            tables=tables,
            metadata={"parser": "python-docx"},
        )

    def _parse_text(self, file_path: Path) -> ParsedDocument:
        """Markdown / TXT 直接读取"""
        text = file_path.read_text(encoding="utf-8")
        return ParsedDocument(
            source_file=file_path.name,
            doc_type=self._infer_doc_type(file_path.name),
            title=file_path.stem,
            text=text,
            metadata={"parser": "direct"},
        )

    def _parse_tabular(self, file_path: Path) -> ParsedDocument:
        """Excel / CSV → 结构化文本块"""
        import pandas as pd

        if file_path.suffix.lower() == ".csv":
            df = pd.read_csv(file_path, encoding="utf-8-sig")
        else:
            df = pd.read_excel(file_path)

        # 将每行转为描述性文本
        lines = []
        for idx, row in df.iterrows():
            parts = [f"{col}: {val}" for col, val in row.items() if pd.notna(val)]
            lines.append(f"[行{idx+1}] " + ", ".join(parts))

        return ParsedDocument(
            source_file=file_path.name,
            doc_type="结构化数据",
            title=file_path.stem,
            text="\n".join(lines),
            tables=[{"rows": df.values.tolist(), "columns": df.columns.tolist()}],
            metadata={"parser": "pandas", "row_count": len(df)},
        )

    def save_parsed(self, docs: list[ParsedDocument]):
        """将解析结果持久化为 JSON，供后续管线读取"""
        for doc in docs:
            out_path = self.parsed_dir / f"{Path(doc.source_file).stem}.json"
            data = {
                "source_file": doc.source_file,
                "doc_type": doc.doc_type,
                "title": doc.title,
                "text": doc.text,
                "tables": doc.tables,
                "figures": doc.figures,
                "page_count": doc.page_count,
                "metadata": doc.metadata,
            }
            out_path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
            logger.info(f"保存解析结果: {out_path}")

    def load_parsed(self) -> list[ParsedDocument]:
        """从 parsed_dir 加载已解析文档"""
        docs = []
        for json_file in sorted(self.parsed_dir.glob("*.json")):
            data = json.loads(json_file.read_text(encoding="utf-8"))
            docs.append(ParsedDocument(**data))
        return docs

    @staticmethod
    def _infer_doc_type(filename: str) -> str:
        """根据文件名推断文档类型"""
        lower = filename.lower()
        if any(kw in lower for kw in ("报告", "勘查", "普查", "详查", "勘探")):
            return "勘查报告"
        elif any(kw in lower for kw in ("图", "geological_map", "地质图")):
            return "地质图"
        elif any(kw in lower for kw in ("论文", "article", "paper", "journal")):
            return "论文"
        else:
            return "技术文档"
